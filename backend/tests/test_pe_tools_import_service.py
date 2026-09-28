"""Parsing des CSV PE Tools ("PeTool - 7.<CODE>.csv") vers raw_data.pe_tools.

Format constate sur les exports du 2026-07-06 : cp850, separateur ';', en-tetes
metier accentues ("Désignation", "Date rév."), guillemets parfois non fermes,
lignes parfois plus larges que l'en-tete. Le service est pur (sans Flask ni
base) pour etre teste ici tel quel.
"""
import os

import pytest

from services.pe_tools_import_service import (
    PE_TOOLS_COLUMNS,
    cle,
    parse_pe_tools_csv,
    parse_pe_tools_excel,
)

FIXTURE = os.path.join(os.path.dirname(__file__), 'fixtures', 'petool_mcar_extrait.csv')

ENTETE = ('Localisation / Classement;Gamme en DMS;Poste technique;Niveau SAP;Plan Entretien;'
          'Poste entretien;Groupe de Gamme;Compteur de Gamme;Frequence;Désignation;Type;Criticité;'
          'Parité semaine;Jour;Décal.;Date de validation;Lien Fichier de gamme Source;'
          'Lien Fichier DMS SAP en PDF;DMS_SAP;Charge;Nombre intervenants;Date rév.;'
          'Nb jours depuis la dernière rév.')


def _cp850(texte: str) -> bytes:
    return texte.encode('cp850')


def test_cle_normalise_accents_casse_espaces_et_ponctuation_finale():
    assert cle('Date rév.') == 'date rev'
    assert cle('DATE  REV') == 'date rev'
    assert cle('Désignation ') == 'designation'
    assert cle('Nb jours depuis la dernière rév.') == 'nb jours depuis la derniere rev'


def test_les_23_colonnes_sont_declarees_dans_l_ordre_de_la_table():
    assert len(PE_TOOLS_COLUMNS) == 23
    assert PE_TOOLS_COLUMNS[0] == ('localisation_classement', 'Localisation / Classement')
    assert PE_TOOLS_COLUMNS[-1] == ('nb_jours_depuis_derniere_rev', 'Nb jours depuis la dernière rév.')


def test_fichier_reel_cp850_mappe_les_23_colonnes():
    with open(FIXTURE, 'rb') as f:
        parsed = parse_pe_tools_csv(f.read())
    assert parsed.missing_columns == []
    assert parsed.unknown_columns == []
    assert parsed.repaired_lines == 0
    assert len(parsed.rows) == 2
    premiere = parsed.rows[0]
    assert set(premiere) == {c for c, _ in PE_TOOLS_COLUMNS}
    assert premiere['poste_technique'] == 'T120-L020'
    assert premiere['plan_entretien'] == '34125'
    assert premiere['designation'] == 'PREV 12S 2MEx4 ON CT CAPTEURS MSA2'
    assert premiere['criticite'] is None          # champ vide -> NULL
    assert premiere['nb_jours_depuis_derniere_rev'] == '663'


def test_utf8_bom_est_accepte():
    contenu = ('﻿' + ENTETE + '\r\n' + 'Voie ferrée;OUI;T410-C;RESEAU;35553;70453;523240;1;4S;'
               'P/4S CONTRÔLE RAILS;CTRL MEC;;;;O;44952;;;;2;2;;\r\n').encode('utf-8')
    parsed = parse_pe_tools_csv(contenu)
    assert parsed.missing_columns == []
    assert parsed.rows[0]['localisation_classement'] == 'Voie ferrée'
    assert parsed.rows[0]['designation'] == 'P/4S CONTRÔLE RAILS'


def test_utf8_sans_bom_est_accepte():
    contenu = (ENTETE + '\r\n' + 'Voie ferrée;OUI;T410-C;RESEAU;35553;70453;523240;1;4S;'
               'P/4S CONTRÔLE RAILS;CTRL MEC;;;;O;44952;;;;2;2;;\r\n').encode('utf-8')
    parsed = parse_pe_tools_csv(contenu)
    assert parsed.missing_columns == []
    assert parsed.rows[0]['localisation_classement'] == 'Voie ferrée'
    assert parsed.rows[0]['designation'] == 'P/4S CONTRÔLE RAILS'


def test_cp1252_excel_windows_est_accepte():
    contenu = (ENTETE + '\r\n' + 'Voie ferrée;OUI;T410-C;RESEAU;35553;70453;523240;1;4S;'
               'P/4S CONTRÔLE RAILS;CTRL MEC;;;;O;44952;;;;2;2;;\r\n').encode('cp1252')
    parsed = parse_pe_tools_csv(contenu)
    assert parsed.missing_columns == []
    assert parsed.rows[0]['localisation_classement'] == 'Voie ferrée'
    assert parsed.rows[0]['designation'] == 'P/4S CONTRÔLE RAILS'


def test_export_de_l_ecran_avec_noms_sql_est_reimportable():
    entete_sql = ';'.join(c for c, _ in PE_TOOLS_COLUMNS)
    contenu = ('﻿' + entete_sql + '\r\n' + 'Voie ferrée;OUI;T410-C;RESEAU;35553;70453;523240;1;4S;'
               'P/4S CONTRÔLE RAILS;CTRL MEC;;;;O;44952;;;;2;2;;\r\n').encode('utf-8')
    parsed = parse_pe_tools_csv(contenu)
    assert parsed.missing_columns == []
    assert parsed.rows[0]['poste_technique'] == 'T410-C'


def test_colonne_manquante_et_colonne_inconnue_sont_signalees():
    entete = ENTETE.replace('Décal.;', '') + ';Commentaire libre'
    contenu = _cp850(entete + '\r\n' + 'A;OUI;T1;N;1;2;3;1;4S;D;T;;;;N;;;;1;1;;;;blabla\r\n')
    parsed = parse_pe_tools_csv(contenu)
    assert parsed.missing_columns == ['Décal.']
    assert parsed.unknown_columns == ['Commentaire libre']
    assert parsed.rows[0]['decalage'] is None
    assert 'Commentaire libre' not in parsed.rows[0]


def test_guillemet_non_ferme_est_repare_et_ligne_trop_large_tronquee():
    contenu = _cp850(
        ENTETE + '\r\n'
        + 'A;OUI;T1;Niveau 12" pouces;1;2;3;1;4S;D;T;;;;N;;;;1;1;;;663\r\n'   # 1 seul guillemet -> repare
        + 'B;OUI;T2;N;1;2;3;1;4S;D;T;;;;N;;;;1;1;;;;surplus1;surplus2\r\n'
    )
    parsed = parse_pe_tools_csv(contenu)
    assert parsed.repaired_lines == 1
    assert len(parsed.rows) == 2
    assert parsed.rows[0]['niveau_sap'] == 'Niveau 12" pouces'
    assert parsed.rows[0]['nb_jours_depuis_derniere_rev'] == '663'
    assert parsed.rows[1]['poste_technique'] == 'T2'
    assert parsed.rows[1]['nb_jours_depuis_derniere_rev'] is None


def test_lignes_vides_ignorees():
    contenu = _cp850(ENTETE + '\r\n;;;;;;;;;;;;;;;;;;;;;;\r\n\r\nA;OUI;T1;N;1;2;3;1;4S;D;T;;;;N;;;;1;1;;;\r\n')
    parsed = parse_pe_tools_csv(contenu)
    assert len(parsed.rows) == 1


def test_contenu_vide_refuse():
    with pytest.raises(ValueError, match='vide'):
        parse_pe_tools_csv(b'')


def test_en_tete_hors_format_refusee():
    contenu = _cp850('col1;col2;col3\r\n1;2;3\r\n')
    with pytest.raises(ValueError, match='Poste technique'):
        parse_pe_tools_csv(contenu)


def test_excel_onglet_7_code_entete_ligne_2_lignes_modele_ignorees():
    """Classeur d'origine : bandeau en ligne 1, en-tete en ligne 2, ~1 000
    lignes de modele sans poste ni plan (ignorees), colonnes hors pe_tools."""
    import datetime
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = 'Aide'
    ws = wb.create_sheet('7.MENG')
    ws.append([None, 'CARTOGRAPHIE DES PE'])
    ws.append(['Poste Responsable', "Poste technique ou N° d'équipement", 'Plan Entretien',
               'Poste entretien', 'Frequence', 'Date de validation', 'Charge', 'Décal.'])
    ws.append(['7.MENG', 'T410-C', 35553, 70453, '4S', datetime.datetime(2023, 1, 26), 1.75, 'O'])
    ws.append(['7.MENG', None, None, None, None, None, None, 'N'])    # ligne de modele
    ws.append(['7.MENG', 'T410-D', 35554, None, '#N/A', None, 2.0, None])
    tampon = io.BytesIO()
    wb.save(tampon)

    res = parse_pe_tools_excel(tampon.getvalue())

    assert len(res.rows) == 2
    r = res.rows[0]
    assert (r['poste_technique'], r['plan_entretien'], r['poste_entretien']) == ('T410-C', '35553', '70453')
    assert (r['date_validation'], r['charge'], r['decalage']) == ('26/01/2023', '1,75', 'O')
    assert (res.rows[1]['frequence'], res.rows[1]['charge']) == (None, '2')
    assert 'Poste Responsable' in res.unknown_columns


def test_excel_sans_onglet_pe_tools_refuse():
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["Plan d'entretien", 'Date dernière exécution'])
    tampon = io.BytesIO()
    wb.save(tampon)

    with pytest.raises(ValueError, match='Plan Entretien'):
        parse_pe_tools_excel(tampon.getvalue())
