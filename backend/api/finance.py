"""Menu Finance : ecrans Immobilisations (clean_data.immobilisation),
Commandes d'achat (clean_data.commande_achat_ifs) et Comptes (conversion des
comptes PCG SAP -> IFS, transcodification FA_ACCOUNT).

Pour chaque ecran <e> = immobilisations | commandes-achat :
  GET  /api/v1/finance/<e>              liste paginee + filtres + stats + definitions de colonnes
  GET  /api/v1/finance/<e>/export.xlsx  classeur Excel (memes filtres que la liste)
  POST /api/v1/finance/<e>/sync         rechargement de la table clean_data ; corps
                                         {"source": "sap"} (defaut : re-extraction SAP
                                         puis rechargement) ou {"source": "mf"}
                                         (rechargement seul depuis raw_data : applique
                                         les transcodifications / valeurs par defaut)
  GET  /api/v1/finance/<e>/sync         etat de la derniere synchronisation
Immobilisations :
  ?vue=immobilisations (defaut)         les 60 colonnes du fichier livre, toutes les fiches
  ?vue=travail                          onglet « Travail » du classeur metier : fiches a
                                         reprendre, date de sortie effacee, tri par date
                                         d'acquisition, OBJECT_GROUP_ID, comptes / site IFS, OTP
  GET  /api/v1/finance/immobilisations/synthese  KPI + sous-totaux, memes filtres
  GET  /api/v1/finance/comptes          conversion des comptes (FA_ACCOUNT) + comptes des fiches
"""
import io
import json
import logging
import os
import re
import tempfile
import threading
import time
from datetime import date, datetime
from decimal import Decimal

from flask import Blueprint, current_app, jsonify, request, send_file
from flask_jwt_extended import get_jwt_identity, jwt_required
from sqlalchemy import text

from models import db

finance_blueprint = Blueprint('finance', __name__)
logger = logging.getLogger(__name__)

# Tables lues par clean_data.alimenter_immobilisation
# (sql/immobilisation/02_alimenter_immobilisation.sql).
# Uniquement des tables transparentes : une vue SAP s'extrait vide.
IMMO_SAP_TABLES = ['ANLA', 'ANLB', 'ANLC', 'ANLZ', 'ANKT', 'T001', 'T095', 'T095T',
                   'T090NAT', 'CSKT', 'TGSBT', 'PRPS']

EXTRACTION_TIMEOUT_SECONDS = 3600
EXTRACTION_POLL_SECONDS = 10


# gunicorn -w4 : l'etat est publie dans un fichier partage par les workers du
# conteneur (meme principe que api/etl.py).
# ponytail: un fichier par ecran = une seule synchro a la fois par ecran ; perdu au
# redeploiement (le bouton redevient simplement disponible).
def _status_file(key):
    return os.path.join(tempfile.gettempdir(), f'{key}_sync.json')


# (colonne, libelle de l'extraction transmise aux metiers) : en-tetes de l'onglet
# « Immobilisations » du classeur, repris tels quels dans l'ecran et l'Excel.
LIBELLES = [
    ('societe_sap', 'Société SAP (ANLA-BUKRS)'),
    ('num_immobilisation', 'Numéro immobilisation (ANLA-ANLN1)'),
    ('sous_numero', 'Sous-numéro immobilisation (ANLA-ANLN2)'),
    ('cle_immobilisation', 'Clé immobilisation (ANLA-BUKRS/ANLN1/ANLN2)'),
    ('libelle', 'Libellé immobilisation (ANLA-TXT50)'),
    ('libelle_complementaire', 'Libellé immobilisation complémentaire (ANLA-TXA50)'),
    ('classe_immo', 'Classe immobilisation (ANLA-ANLKL)'),
    ('famille_immo', 'Famille immobilisation (ANKT)'),
    ('indicateur_suppression', 'Indicateur suppression (ANLA-XLOEV)'),
    ('numero_serie', 'Numéro de série (ANLA-SERNR)'),
    ('pays', 'Pays (ANLA-LAND1)'),
    ('groupe_evaluation_1', 'Groupe évaluation 1 (ANLA-ORD41)'),
    ('groupe_evaluation_2', 'Groupe évaluation 2 (ANLA-ORD42)'),
    ('groupe_evaluation_3', 'Groupe évaluation 3 (ANLA-ORD43)'),
    ('groupe_evaluation_4', 'Groupe évaluation 4 (ANLA-ORD44)'),
    ('projet', 'Projet (ANLA-PROJN)'),
    ('cle_comptes_immo', 'Clé détermination comptes immo SAP (ANLA-KTOGR)'),
    ('libelle_cle_comptes_immo', 'Libellé clé comptable immo SAP (T095T-KTGRTX)'),
    ('compte_immobilisation', 'Compte immobilisation SAP (T095-KTANSW, AFABE 02)'),
    ('compte_amort_cumule', 'Compte amortissement cumulé SAP (T095-KTANZA, AFABE 02)'),
    ('compte_dotation_amort', 'Compte dotation amortissement SAP (T095-KTAUFG, AFABE 02)'),
    ('date_acquisition', 'Date acquisition / capitalisation (ANLA-AKTIV)'),
    ('date_premiere_acquisition', 'Date première acquisition SAP (ANLA-ZUGDT)'),
    ('date_debut_amort', 'Date début amortissement (ANLB-AFABG)'),
    ('date_fin_amort_estimee', 'Date fin amortissement estimée (calcul ANLB/ANLC)'),
    ('duree_amort_annees', 'Durée amortissement années (ANLB-NDJAR)'),
    ('duree_amort_periodes', 'Durée amortissement mois / périodes (ANLB-NDPER)'),
    ('duree_amort_totale_mois', 'Durée amortissement totale mois (calcul ANLB)'),
    ('zone_amortissement', 'Zone amortissement retenue (ANLB-AFABE)'),
    ('type_amortissement', 'Type amortissement SAP (ANLB-AFASL)'),
    ('libelle_type_amortissement', 'Libellé type amortissement SAP (T090NAT-AFATXT)'),
    ('taux_amort_estime', 'Taux amortissement estimé % (calcul)'),
    ('centre_cout', 'Centre de coût amortissements (ANLZ-KOSTL)'),
    ('libelle_centre_cout', 'Libellé centre de coût (CSKT)'),
    ('site_sap', 'Site SAP (ANLZ-WERKS)'),
    ('secteur_sap', 'Secteur SAP (ANLZ-GSBER)'),
    ('libelle_secteur', 'Libellé secteur SAP (TGSBT)'),
    ('emplacement', 'Lieu / emplacement SAP (ANLZ-STORT)'),
    ('immo_origine', "Immobilisation d'origine principale (ANLA-AIBN1)"),
    ('sous_numero_origine', "Sous-numéro immobilisation d'origine (ANLA-AIBN2)"),
    ('date_origine', 'Date origine immobilisation (ANLA-AIBDT)'),
    ('numero_inventaire', 'Numéro inventaire (ANLA-INVNR)'),
    ('fabricant', 'Fabricant (ANLA-HERST)'),
    ('type_modele', 'Type / modèle (ANLA-TYPBZ)'),
    ('fournisseur', 'Fournisseur (ANLA-LIFNR)'),
    ('quantite', 'Quantité (ANLA-MENGE)'),
    ('unite', 'Unité (ANLA-MEINS)'),
    ('ordre_investissement', 'Ordre investissement (ANLA-EAUFN)'),
    ('zone_valorisation', 'Zone valorisation retenue (ANLC-AFABE)'),
    ('exercice_valorisation', 'Exercice de valorisation retenu (ANLC-GJAHR)'),
    ('valeur_acq_debut_exercice', 'Valeur acquisition début exercice (ANLC-KANSW)'),
    ('mouvements_acq_exercice', 'Mouvements acquisition exercice (ANLC-ANSWL)'),
    ('sorties_exercice', 'Sorties valeur exercice (ANLC-ABGAN)'),
    ('valeur_acq_fin_exercice', 'Valeur acquisition fin exercice (ANLC-KANSW + ANLC-ANSWL - ANLC-ABGAN)'),
    ('amort_cumules', "Amortissements cumulés à l'ouverture de l'exercice (ANLC-KNAFA/KSAFA/KAAFA/KMAFA)"),
    ('vnc', "VNC à l'ouverture de l'exercice (calcul)"),
    ('dotation_annuelle', "Dotation annuelle comptabilisée (non incluse à l'ouverture)"),
    ('blocage_comptabilisation', 'Blocage comptabilisation (ANLA-XSPEB)'),
    ('date_sortie', 'Date sortie (ANLA-ABGDT)'),
    ('date_desactivation', 'Date désactivation (ANLA-DEAKT)'),
    # Reprise IFS (migration 098, classeur metier du 18/08/2026)
    ('reprise_ifs', 'Reprise IFS'),
    ('motif_exclusion', "Motif d'exclusion"),
    ('compte_immobilisation_ifs', 'Compte immobilisation IFS'),
    ('compte_amort_cumule_ifs', 'Compte amortissement cumulé IFS'),
    ('object_group_id', 'OBJECT_GROUP_ID'),
    ('site_ifs', 'Site IFS'),
    ('element_otp', "Élément d'OTP (ANLA-POSNR -> PRPS-POSID)"),
    ('libelle_otp', 'Libellé OTP (PRPS-POST1)'),
]
LIBELLE = dict(LIBELLES)
COLONNES = [c for c, _ in LIBELLES]
_MONTANTS = {'valeur_acq_debut_exercice', 'mouvements_acq_exercice', 'sorties_exercice',
             'valeur_acq_fin_exercice', 'amort_cumules', 'vnc', 'dotation_annuelle'}

# Vue « immobilisations » = les 60 colonnes du fichier livre.
COLONNES_IMMO = COLONNES[:60]
# Vue « travail » = onglet Travail du classeur : les 60 colonnes avec OBJECT_GROUP_ID
# insere apres la duree en annees (meme place que dans le classeur), puis les
# colonnes converties (comptes IFS, site, OTP) que la Methode demande de remplir.
_i = COLONNES_IMMO.index('duree_amort_annees') + 1
COLONNES_TRAVAIL = (COLONNES_IMMO[:_i] + ['object_group_id'] + COLONNES_IMMO[_i:]
                    + ['compte_immobilisation_ifs', 'compte_amort_cumule_ifs', 'site_ifs',
                       'element_otp', 'libelle_otp'])

# Etapes de l'onglet « Méthode » du classeur et leur automatisation (affichees
# dans l'onglet Travail de l'ecran).
METHODE = [
    ("1/ Vérifier dans l'extraction SAP le cumul des comptes immo brutes et le cumul des comptes "
     "d'amortissement par rapport à la balance générale.",
     "Synthèse > sous-totaux par compte (acquisition / amortissements) ; le rapprochement avec la balance reste manuel."),
    ("2/ Supprimer les immobilisations ayant une date de sortie antérieure ou égale à la date de bascule "
     "IFS ; reprendre celles sorties après la bascule en effaçant leur date de sortie.",
     "Colonne Reprise IFS (date de bascule = Valeurs par défaut) ; l'onglet Travail ne contient que les "
     "fiches à reprendre, date de sortie effacée."),
    ("3/ Convertir au format date les dates de désactivation.", "Dates typées au chargement."),
    ("4/ Recontrôler les totaux (valeurs brutes et amortissements), qui doivent rester identiques.",
     "Synthèse sur la vue Immobilisations et sur la vue Travail."),
    ("5/ Retrier les immobilisations restantes par date d'acquisition.", "Tri de l'onglet Travail."),
    ("6/ Mettre à blanc les dates en 00/01/1900.", "Dates SAP 00000000 chargées à NULL."),
    ("7/ Remplir les colonnes du gabarit : groupe objet, comptes PCG -> IFS, site SJ / CS.",
     "OBJECT_GROUP_ID (Transcodification FA_OBJECT_GROUP / FA_OBJECT_GROUP_IMMO), comptes IFS "
     "(FA_ACCOUNT, écran Finance > Comptes), Site IFS (division 9200 -> SJ, 9000 -> CS)."),
]


def _kind(col):
    if col == 'reprise_ifs':
        return 'bool'
    if col.startswith('date_'):
        return 'date'
    if col in _MONTANTS:
        return 'montant'
    if col in ('duree_amort_totale_mois', 'taux_amort_estime'):
        return 'nombre'
    return 'text'


def _defs(cols):
    return [{'key': c, 'label': LIBELLE[c], 'kind': _kind(c)} for c in cols]


def _read_status(key):
    try:
        with open(_status_file(key), encoding='utf-8') as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def _write_status(key, **fields):
    status = {**(_read_status(key) or {}), **fields}
    tmp = _status_file(key) + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(status, f, default=str)
    os.replace(tmp, _status_file(key))
    return status


def _json(v):
    """date -> 'YYYY-MM-DD', Decimal -> float (jsonify ferait une date HTTP et une chaine)."""
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    return v


def _now():
    return datetime.now().isoformat(timespec='seconds')


# ---------------------------------------------------------------------------
# Immobilisations
# ---------------------------------------------------------------------------

def _vue():
    return 'travail' if request.args.get('vue') == 'travail' else 'immobilisations'


def _filtres(vue=None):
    """Clause WHERE + parametres a partir de la query string (liste, synthese, Excel)."""
    vue = vue or _vue()
    search = (request.args.get('search') or '').strip()
    secteur = (request.args.get('secteur') or '').strip()
    statut = (request.args.get('statut') or '').strip()
    reprise = 'a_reprendre' if vue == 'travail' else (request.args.get('reprise') or '').strip()

    where, params = [], {}
    if reprise == 'a_reprendre':
        where.append('reprise_ifs')
    elif reprise == 'exclues':
        where.append('NOT reprise_ifs')
    if search:
        where.append("(num_immobilisation ILIKE :s OR libelle ILIKE :s OR famille_immo ILIKE :s "
                     "OR centre_cout ILIKE :s OR numero_inventaire ILIKE :s "
                     "OR compte_immobilisation ILIKE :s)")
        params['s'] = f'%{search}%'
    if secteur:
        where.append('secteur_sap = :secteur')
        params['secteur'] = secteur
    if statut == 'actives':
        where.append('date_desactivation IS NULL')
    elif statut == 'desactivees':
        where.append('date_desactivation IS NOT NULL')
    clause = ('WHERE ' + ' AND '.join(where)) if where else ''
    return clause, params


def _select_immo(vue, clause, params, limit=None):
    """Lignes d'une vue. Travail : date de sortie effacee (etape 2 de la Methode),
    tri par date d'acquisition (etape 5)."""
    cols = COLONNES_TRAVAIL if vue == 'travail' else COLONNES_IMMO
    exprs = ['NULL::date AS date_sortie' if vue == 'travail' and c == 'date_sortie' else c
             for c in cols]
    order = ('date_acquisition NULLS LAST, num_immobilisation, sous_numero' if vue == 'travail'
             else 'num_immobilisation, sous_numero')
    sql = f'SELECT {", ".join(exprs)} FROM clean_data.immobilisation {clause} ORDER BY {order}'
    if limit:
        sql += ' LIMIT :limit OFFSET :offset'
    return cols, db.session.execute(text(sql), params).mappings().all()


@finance_blueprint.route('/immobilisations', methods=['GET'])
@jwt_required()
def list_immobilisations():
    vue = _vue()
    page = max(int(request.args.get('page', 1)), 1)
    page_size = min(max(int(request.args.get('page_size', 50)), 1), 500)
    clause, params = _filtres(vue)

    total = db.session.execute(
        text(f'SELECT count(*) FROM clean_data.immobilisation {clause}'), params).scalar()
    cols, rows = _select_immo(vue, clause, {**params, 'limit': page_size,
                                            'offset': (page - 1) * page_size}, limit=True)

    # Totaux sur la selection filtree (valeurs statutaires a l'ouverture)
    stats = db.session.execute(text(
        'SELECT count(*) AS immobilisations, '
        'count(*) FILTER (WHERE date_desactivation IS NULL) AS actives, '
        'count(*) FILTER (WHERE reprise_ifs) AS a_reprendre, '
        'count(*) FILTER (WHERE NOT reprise_ifs) AS exclues, '
        'count(*) FILTER (WHERE reprise_ifs AND compte_immobilisation IS NOT NULL '
        '                   AND compte_immobilisation_ifs IS NULL) AS sans_compte_ifs, '
        'count(*) FILTER (WHERE reprise_ifs AND object_group_id IS NULL) AS sans_groupe_objet, '
        'sum(valeur_acq_debut_exercice) AS valeur_acquisition, '
        'sum(amort_cumules) AS amort_cumules, sum(vnc) AS vnc, '
        'max(exercice_valorisation) AS exercice, '
        "public.get_default_value('clean_data.immobilisation', 'date_bascule_ifs', 'STANDARD') AS date_bascule "
        f'FROM clean_data.immobilisation {clause}'), params).mappings().one()
    secteurs = [dict(r) for r in db.session.execute(text(
        'SELECT secteur_sap AS code, max(libelle_secteur) AS libelle '
        'FROM clean_data.immobilisation WHERE secteur_sap IS NOT NULL '
        'GROUP BY secteur_sap ORDER BY secteur_sap')).mappings()]

    return jsonify({
        'vue': vue,
        'colonnes': _defs(cols),
        'methode': METHODE if vue == 'travail' else None,
        'rows': [{k: _json(v) for k, v in r.items()} for r in rows],
        'total': total,
        'page': page,
        'page_size': page_size,
        'stats': {k: _json(v) for k, v in stats.items()},
        'secteurs': secteurs,
        'sync': _read_status('immobilisation'),
    })


# Axes de la synthese : (cle json, colonne, titre de bloc) — memes blocs que
# l'onglet « Synthèse » du classeur metier, plus le compte pour le controle des
# cumuls par compte general (etapes 1 et 4 de la Methode).
_SYNTHESE_AXES = [
    ('par_type', 'libelle_type_amortissement', "Sous-totaux par type d'amortissement"),
    ('par_famille', 'famille_immo', 'Sous-totaux par famille'),
    ('par_secteur', 'libelle_secteur', 'Sous-totaux par secteur'),
    ('par_compte', 'compte_immobilisation', 'Sous-totaux par compte immobilisation (SAP -> IFS)'),
]


def _synthese(clause, params):
    """KPI + sous-totaux en une requete (GROUPING SETS), sur la selection filtree."""
    axes = [col for _, col, _ in _SYNTHESE_AXES]
    rows = db.session.execute(text(
        'SELECT ' + ', '.join(f'GROUPING({c}) AS g_{c}' for c in axes) + ', '
        + ', '.join(axes) + ', max(compte_immobilisation_ifs) AS compte_ifs, '
        'count(*) AS nombre, sum(valeur_acq_debut_exercice) AS acquisition, '
        'sum(amort_cumules) AS amortissements, sum(vnc) AS vnc, '
        'max(exercice_valorisation) AS exercice '
        f'FROM clean_data.immobilisation {clause} '
        'GROUP BY GROUPING SETS ((), ' + ', '.join(f'({c})' for c in axes) + ')'
    ), params).mappings().all()

    result = {'kpi': {}, **{k: [] for k, _, _ in _SYNTHESE_AXES}}
    for row in rows:
        mesures = {k: _json(row[k]) for k in ('nombre', 'acquisition', 'amortissements', 'vnc')}
        actifs = [c for c in axes if not row[f'g_{c}']]
        if not actifs:
            exercice = row['exercice']
            result['kpi'] = {
                **mesures, 'exercice': exercice,
                'date_situation': f'01/07/{int(exercice) - 1}' if exercice else None,
                'zone': '02 - Amortissement statutaire CRC2002-10',
            }
            continue
        key = next(k for k, c, _ in _SYNTHESE_AXES if c == actifs[0])
        libelle = row[actifs[0]]
        if actifs[0] == 'compte_immobilisation' and libelle:
            libelle = f"{libelle} -> {row['compte_ifs'] or '(sans compte IFS)'}"
        result[key].append({'libelle': libelle or '(vide)', **mesures})
    for k, _, _ in _SYNTHESE_AXES:
        result[k].sort(key=lambda x: x['libelle'])
    return result


@finance_blueprint.route('/immobilisations/synthese', methods=['GET'])
@jwt_required()
def synthese_immobilisations():
    clause, params = _filtres()
    return jsonify(_synthese(clause, params))


@finance_blueprint.route('/immobilisations/export.xlsx', methods=['GET'])
@jwt_required()
def export_immobilisations_excel():
    """Classeur Excel a l'image du classeur metier : onglets Immobilisations (60
    colonnes), Travail, Conversion cpte général, TRansco comptes generaux et
    Synthèse. Les filtres de l'ecran (recherche, secteur, statut, reprise)
    s'appliquent a l'onglet Immobilisations et a la Synthese ; Travail ne
    retient que les fiches a reprendre."""
    from openpyxl import Workbook

    clause, params = _filtres('immobilisations')
    clause_t, params_t = _filtres('travail')
    wb = Workbook()
    # _select_immo renvoie des mappings : iterer une ligne donnerait les noms de colonnes
    cols, rows = _select_immo('immobilisations', clause, params)
    _feuille_tableau(wb.active, 'Immobilisations', [(c, LIBELLE[c]) for c in cols],
                     [[r[c] for c in cols] for r in rows], _MONTANTS)
    cols, rows = _select_immo('travail', clause_t, params_t)
    _feuille_tableau(wb.create_sheet('Travail'), 'Travail', [(c, LIBELLE[c]) for c in cols],
                     [[r[c] for c in cols] for r in rows], _MONTANTS)

    conv = _conversion_comptes()
    _feuille_tableau(wb.create_sheet('Conversion cpte général'), 'Conversion cpte général',
                     [('compte_sap', 'PCG US SAP'), ('pcg_fr', 'PCG FR SAP'),
                      ('compte_ifs', 'PCG IFS'), ('libelle', 'LIBELLE')],
                     [(c['compte_sap'], c['pcg_fr'], c['compte_ifs'], c['libelle']) for c in conv], set())
    transco = db.session.execute(text(
        'SELECT num_immobilisation, compte_immobilisation, compte_immobilisation_ifs '
        f'FROM clean_data.immobilisation {clause_t} '
        'ORDER BY date_acquisition NULLS LAST, num_immobilisation, sous_numero'), params_t).all()
    _feuille_tableau(wb.create_sheet('TRansco comptes generaux'), 'TRansco comptes generaux',
                     [('num', LIBELLE['num_immobilisation']), ('cpt', 'Cpt fiche'),
                      ('nv', 'Nv cpte fiche')], transco, set())
    synthese = _synthese(clause, params)
    # Lectures terminees : liberer la transaction (et ses verrous) avant de
    # construire le classeur, qui prend plusieurs secondes.
    db.session.commit()
    _feuille_synthese(wb.create_sheet('Synthèse'), synthese)
    return _envoyer(wb, 'immobilisations')


def _sync_source():
    """'sap' (defaut : extraction puis rechargement) ou 'mf' (rechargement seul)."""
    body = request.get_json(silent=True) or {}
    return 'mf' if body.get('source') == 'mf' else 'sap'


@finance_blueprint.route('/immobilisations/sync', methods=['GET'])
@jwt_required()
def get_sync_status():
    return jsonify(_read_status('immobilisation') or {'status': 'never'})


@finance_blueprint.route('/immobilisations/sync', methods=['POST'])
@jwt_required()
def start_sync():
    return _start_sync('immobilisation', IMMO_SAP_TABLES, 'clean_data.immobilisation',
                       lambda conn: conn.execute(
                           text('SELECT clean_data.alimenter_immobilisation()')).scalar(),
                       extraire=_sync_source() == 'sap')


# ---------------------------------------------------------------------------
# Comptes : conversion PCG SAP -> IFS (transcodification FA_ACCOUNT) et comptes
# portes par les fiches (onglets « Conversion cpte général » et « TRansco
# comptes generaux » du classeur)
# ---------------------------------------------------------------------------

# La description d'une ligne FA_ACCOUNT porte « LIBELLE (PCG FR nnn) » : le PCG
# francais n'a pas de colonne propre dans TranscodificationTable.
_PCG_FR = re.compile(r'\s*\(PCG FR ([^)]*)\)\s*$')


def _conversion_comptes():
    rows = db.session.execute(text(
        'SELECT id, source_value, target_value, description, is_active, updated_at '
        'FROM public."TranscodificationTable" '
        "WHERE category = 'FA_ACCOUNT' AND source_system = 'SAP' AND target_system = 'IFS' "
        'ORDER BY source_value')).mappings().all()
    out = []
    for r in rows:
        desc = r['description'] or ''
        m = _PCG_FR.search(desc)
        out.append({
            'id': r['id'], 'compte_sap': r['source_value'], 'compte_ifs': r['target_value'],
            'pcg_fr': m.group(1) if m else None,
            'libelle': _PCG_FR.sub('', desc).strip(),
            'is_active': r['is_active'], 'updated_at': _json(r['updated_at']),
        })
    return out


@finance_blueprint.route('/comptes', methods=['GET'])
@jwt_required()
def comptes():
    """Conversion des comptes + comptes reellement portes par les fiches, avec
    la conversion resolue EN DIRECT (une saisie est visible sans recharger ;
    les colonnes de la table, elles, suivent au prochain rechargement)."""
    fiches = db.session.execute(text(
        "SELECT role, compte_sap, "
        "       public.get_transcodification('FA_ACCOUNT', compte_sap, 'SAP', 'IFS') AS compte_ifs, "
        "       count(*) AS nb_fiches, count(*) FILTER (WHERE reprise_ifs) AS nb_reprises, "
        "       sum(valeur_acq_debut_exercice) AS acquisition, sum(amort_cumules) AS amortissements, "
        "       sum(vnc) AS vnc "
        "FROM (SELECT 'immobilisation' AS role, compte_immobilisation AS compte_sap, reprise_ifs, "
        "             valeur_acq_debut_exercice, amort_cumules, vnc FROM clean_data.immobilisation "
        "      UNION ALL "
        "      SELECT 'amortissement', compte_amort_cumule, reprise_ifs, "
        "             valeur_acq_debut_exercice, amort_cumules, vnc FROM clean_data.immobilisation) u "
        "WHERE compte_sap IS NOT NULL GROUP BY role, compte_sap ORDER BY role, compte_sap"
    )).mappings().all()
    return jsonify({
        'conversion': _conversion_comptes(),
        'fiches': [{k: _json(v) for k, v in r.items()} for r in fiches],
        'sync': _read_status('immobilisation'),
    })


# ---------------------------------------------------------------------------
# Commandes d'achat (clean_data.commande_achat_ifs, module sql/commandeAchat/)
# ---------------------------------------------------------------------------

# Tables lues par clean_data.alimenter_commande_achat_ifs
# (sql/commandeAchat/02_alimenter_commande_achat_ifs.sql).
CA_SAP_TABLES = ['EKKO', 'EKPO', 'EKBE', 'EKET', 'EKPA', 'EKKN', 'LFA1', 'T001W', 'ADRC', 'PRPS']

CA_LIBELLES = [
    ('site', 'Site'),
    ('societe_sap', 'Société SAP'),
    ('num_commande_sap', 'N° commande SAP'),
    ('num_ligne_sap', 'N° ligne SAP'),
    ('fournisseur_sap', 'Fournisseur SAP'),
    ('fournisseur_ifs', 'Fournisseur IFS'),
    ('nom_fournisseur', 'Nom fournisseur'),
    ('fournisseur_facturation_sap', 'Fournisseur facturation SAP'),
    ('fournisseur_facturation_ifs', 'Fournisseur facturation IFS'),
    ('type_ligne_ifs', 'Type ligne IFS'),
    ('article_sap', 'Article SAP'),
    ('designation', 'Désignation'),
    ('qte_commandee', 'Quantité commandée'),
    ('qte_restant_livrer', 'Quantité restant à livrer'),
    ('qte_restant_facturer', 'Quantité restant à facturer'),
    ('unite_achat', "Unité d'achat"),
    ('prix_net_unitaire', 'Prix net unitaire'),
    ('montant_restant_livrer', 'Montant restant à livrer'),
    ('montant_restant_facturer', 'Montant restant à facturer'),
    ('devise', 'Devise'),
    ('taux_change', 'Taux de change'),
    ('date_creation', 'Date création'),
    ('date_livraison_planifiee', 'Date livraison planifiée'),
    ('date_reception_souhaitee', 'Date réception souhaitée'),
    ('date_livraison_promise', 'Date livraison promise'),
    ('acheteur_sap', 'Acheteur SAP'),
    ('condition_paiement', 'Condition de paiement'),
    ('condition_livraison', 'Condition de livraison'),
    ('mode_expedition', "Mode d'expédition"),
    ('adresse_livraison', 'Adresse de livraison'),
    ('code_postal_livraison', 'Code postal livraison'),
    ('ville_livraison', 'Ville livraison'),
    ('pays_livraison', 'Pays livraison'),
    ('pre_imputation_projet', 'Pré-imputation projet'),
]
CA_COLONNES = [c for c, _ in CA_LIBELLES]
_CA_MONTANTS = {'prix_net_unitaire', 'montant_restant_livrer', 'montant_restant_facturer'}
_CA_QUANTITES = {'qte_commandee', 'qte_restant_livrer', 'qte_restant_facturer'}
# Les dates de commande_achat_ifs sont du texte 'JJ/MM/AAAA' (type de la base reelle).
_CA_DATES = {'date_creation', 'date_livraison_planifiee', 'date_reception_souhaitee',
             'date_livraison_promise'}


def _ca_filtres():
    """Clause WHERE + parametres (liste et export Excel des commandes d'achat)."""
    search = (request.args.get('search') or '').strip()
    site = (request.args.get('site') or '').strip()

    where, params = [], {}
    if search:
        where.append("(num_commande_sap ILIKE :s OR nom_fournisseur ILIKE :s "
                     "OR fournisseur_sap ILIKE :s OR fournisseur_ifs ILIKE :s "
                     "OR article_sap ILIKE :s OR designation ILIKE :s)")
        params['s'] = f'%{search}%'
    if site:
        where.append('site = :site')
        params['site'] = site
    clause = ('WHERE ' + ' AND '.join(where)) if where else ''
    return clause, params


def _ca_date(v):
    """'JJ/MM/AAAA' -> date pour Excel (texte laisse tel quel si autre format)."""
    try:
        return datetime.strptime(v, '%d/%m/%Y').date() if v else v
    except (TypeError, ValueError):
        return v


@finance_blueprint.route('/commandes-achat', methods=['GET'])
@jwt_required()
def list_commandes_achat():
    page = max(int(request.args.get('page', 1)), 1)
    page_size = min(max(int(request.args.get('page_size', 50)), 1), 500)
    clause, params = _ca_filtres()

    total = db.session.execute(
        text(f'SELECT count(*) FROM clean_data.commande_achat_ifs {clause}'), params).scalar()
    rows = db.session.execute(text(
        f'SELECT {", ".join(CA_COLONNES)} FROM clean_data.commande_achat_ifs {clause} '
        'ORDER BY num_commande_sap, num_ligne_sap LIMIT :limit OFFSET :offset'
    ), {**params, 'limit': page_size, 'offset': (page - 1) * page_size}).mappings().all()

    # Totaux sur la selection filtree ; montants en EUR seulement (devises non converties)
    stats = db.session.execute(text(
        'SELECT count(*) AS lignes, count(DISTINCT num_commande_sap) AS commandes, '
        'count(DISTINCT fournisseur_sap) AS fournisseurs, '
        "sum(montant_restant_livrer) FILTER (WHERE devise = 'EUR') AS restant_livrer_eur, "
        "sum(montant_restant_facturer) FILTER (WHERE devise = 'EUR') AS restant_facturer_eur "
        f'FROM clean_data.commande_achat_ifs {clause}'), params).mappings().one()
    sites = [r[0] for r in db.session.execute(text(
        'SELECT DISTINCT site FROM clean_data.commande_achat_ifs '
        'WHERE site IS NOT NULL ORDER BY 1'))]

    return jsonify({
        'rows': [{k: _json(v) for k, v in r.items()} for r in rows],
        'total': total,
        'page': page,
        'page_size': page_size,
        'stats': {k: _json(v) for k, v in stats.items()},
        'sites': sites,
        'sync': _read_status('commande_achat'),
    })


@finance_blueprint.route('/commandes-achat/export.xlsx', methods=['GET'])
@jwt_required()
def export_commandes_achat_excel():
    from openpyxl import Workbook

    clause, params = _ca_filtres()
    rows = db.session.execute(text(
        f'SELECT {", ".join(CA_COLONNES)} FROM clean_data.commande_achat_ifs {clause} '
        'ORDER BY num_commande_sap, num_ligne_sap'), params).all()
    dates = [i for i, c in enumerate(CA_COLONNES) if c in _CA_DATES]
    rows = [[_ca_date(v) if i in dates else v for i, v in enumerate(r)] for r in rows]
    db.session.commit()  # liberer la transaction de lecture avant la generation
    wb = Workbook()
    _feuille_tableau(wb.active, "Commandes d'achat", CA_LIBELLES, rows, _CA_MONTANTS | _CA_QUANTITES)
    return _envoyer(wb, 'commandes_achat')


@finance_blueprint.route('/commandes-achat/sync', methods=['GET'])
@jwt_required()
def get_ca_sync_status():
    return jsonify(_read_status('commande_achat') or {'status': 'never'})


@finance_blueprint.route('/commandes-achat/sync', methods=['POST'])
@jwt_required()
def start_ca_sync():
    # Memes bornes que le module ETL : module_params de etl_target_tables
    # ({"date_debut", "date_fin"} sur la date de creation SAP).
    module_params = db.session.execute(text(
        "SELECT module_params FROM public.etl_target_tables "
        "WHERE python_module = 'etl_commande_achat.py' LIMIT 1")).scalar() or {}

    def load(conn):
        return conn.execute(text(
            'SELECT clean_data.alimenter_commande_achat_ifs('
            'CAST(:d AS date), CAST(:f AS date), NULL)'),
            {'d': module_params.get('date_debut'), 'f': module_params.get('date_fin')}).scalar()

    return _start_sync('commande_achat', CA_SAP_TABLES, 'clean_data.commande_achat_ifs', load,
                       extraire=_sync_source() == 'sap')


# ---------------------------------------------------------------------------
# Communs : export Excel, synchronisation SAP
# ---------------------------------------------------------------------------

def _feuille_tableau(ws, title, libelles, rows, numeriques):
    """Feuille tableau : en-tetes metier, dates et nombres types, en-tete fige + filtres."""
    from openpyxl.styles import Alignment, Font, PatternFill

    colonnes = [c for c, _ in libelles]
    ws.title = title[:31]
    ws.append([l for _, l in libelles])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[1].height = 45
    # ws.max_row / ws[n] parcourent toutes les cellules : indices tenus a la main,
    # sinon l'export est quadratique (7 748 x 60 cellules -> ne se termine pas).
    for i, r in enumerate(rows, 2):
        ws.append([float(v) if isinstance(v, Decimal)
                   else ('Oui' if v else 'Non') if isinstance(v, bool) else v
                   for v in r])
        for j in range(1, len(colonnes) + 1):
            cell = ws.cell(i, j)
            # Texte SAP commencant par '=' : openpyxl en ferait une formule -> forcer le texte
            if cell.data_type == 'f':
                cell.data_type = 's'
            elif isinstance(cell.value, (date, datetime)):
                cell.number_format = 'DD/MM/YYYY'

    for idx, col in enumerate(colonnes, 1):
        letter = ws.cell(1, idx).column_letter
        long = col.startswith(('libelle', 'designation', 'nom_', 'adresse'))
        ws.column_dimensions[letter].width = 34 if long else 16
        if col in numeriques:
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = '#,##0.00'
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions


def _feuille_synthese(ws, synthese):
    """Feuille Synthese : en-tete, KPI, un bloc par axe (valeurs, pas de formules)."""
    from openpyxl.styles import Font, PatternFill

    titre = Font(bold=True, color='FFFFFF')
    fond = PatternFill('solid', fgColor='1F4E78')
    sous = PatternFill('solid', fgColor='5B9BD5')
    kpi = synthese['kpi']
    ws['A1'] = 'Synthèse des immobilisations statutaires'
    ws['A1'].font = Font(bold=True, size=14)
    for i, (k, v) in enumerate([('Société', 'STJN'), ('Date de situation', kpi.get('date_situation')),
                                ('Exercice SAP', kpi.get('exercice')), ('Zone retenue', kpi.get('zone'))], 2):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
    ws.cell(7, 1, 'Indicateurs').font = titre
    ws.cell(7, 1).fill = fond
    for i, (k, v) in enumerate([("Nombre d'immobilisations", kpi.get('nombre')),
                                ("Valeur d'acquisition début exercice", kpi.get('acquisition')),
                                ('Amortissements cumulés', kpi.get('amortissements')),
                                ('Valeur nette comptable', kpi.get('vnc'))], 8):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v).number_format = '#,##0.00'
    ligne = 13
    for key, _, bloc in _SYNTHESE_AXES:
        for j, h in enumerate([bloc, 'Nombre', 'Acquisition', 'Amortissements cumulés', 'VNC'], 1):
            c = ws.cell(ligne, j, h)
            c.font = titre
            c.fill = sous
        for item in synthese[key]:
            ligne += 1
            ws.cell(ligne, 1, item['libelle'])
            ws.cell(ligne, 2, item['nombre'])
            for j, k in enumerate(('acquisition', 'amortissements', 'vnc'), 3):
                ws.cell(ligne, j, item[k]).number_format = '#,##0.00'
        ligne += 2
    for col, w in zip('ABCDE', (48, 14, 20, 24, 20)):
        ws.column_dimensions[col].width = w


def _envoyer(wb, prefix):
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True,
        download_name=f'{prefix}_{datetime.now():%Y%m%d_%H%M}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _start_sync(key, tables, cible, load, extraire=True):
    """Lance (extraction SAP +) rechargement dans un thread ; 409 si deja en cours."""
    current = _read_status(key)
    if current and current.get('status') == 'running':
        return jsonify({'error': 'Une synchronisation est deja en cours.', **current}), 409

    user = str(get_jwt_identity() or 'finance-sync')
    status = _write_status(
        key, status='running', step='Demarrage', progress=0, error=None, rows=None,
        source='sap' if extraire else 'mf',
        started_at=_now(), finished_at=None, started_by=user, extraction_id=None,
    )
    app = current_app._get_current_object()
    threading.Thread(target=_run_sync, args=(app, user, key, tables, cible, load, extraire),
                     daemon=True, name=f'{key}-sync').start()
    return jsonify(status), 202


def _run_sync(app, user, key, tables, cible, load, extraire):
    with app.app_context():
        try:
            if extraire:
                _extract(user, key, tables)
            _write_status(key, step=f'Rechargement de {cible}', progress=90)
            with db.engine.begin() as conn:
                rows = load(conn)
            _write_status(key, status='completed', step='Termine', progress=100,
                          rows=rows, finished_at=_now())
        except Exception as e:  # le thread ne doit jamais mourir en silence
            logger.error(f'Synchronisation {key} en echec : {e}')
            _write_status(key, status='failed', error=str(e), finished_at=_now())


def _extract(user, key, tables):
    """Extraction SAP (differentielle) des tables sources, attente de la fin."""
    from services.extraction_service import extraction_service

    _write_status(key, step=f'Extraction SAP ({len(tables)} tables)', progress=5)
    result = extraction_service.start_extraction(
        tables=tables, options={'mode': 'standard'}, user_id=user)
    extraction_id = result.get('extraction_id')
    if not extraction_id:
        raise RuntimeError("Le conteneur d'extraction SAP n'a pas renvoye d'identifiant.")
    _write_status(key, extraction_id=extraction_id)

    deadline = time.time() + EXTRACTION_TIMEOUT_SECONDS
    while time.time() < deadline:
        time.sleep(EXTRACTION_POLL_SECONDS)
        try:
            payload = extraction_service.get_extraction_status(extraction_id)
        except Exception as e:
            logger.warning(f'Statut extraction {extraction_id} indisponible ({e}), nouvel essai.')
            continue
        status = (payload.get('status') or 'running').lower()
        if status in ('completed', 'success', 'done', 'finished'):
            return
        if status in ('failed', 'error'):
            raise RuntimeError(f"L'extraction SAP a echoue : {payload.get('error_message') or 'cause inconnue'}")
        if status in ('stopped', 'cancelled', 'canceled'):
            raise RuntimeError("L'extraction SAP a ete interrompue.")
        progress = payload.get('progress_percentage')
        if progress is not None:
            # L'extraction occupe la plage 5 % -> 85 %.
            _write_status(key, progress=5 + int(float(progress) * 0.8))
    raise RuntimeError(f"Extraction SAP non terminee apres {EXTRACTION_TIMEOUT_SECONDS // 60} min.")
