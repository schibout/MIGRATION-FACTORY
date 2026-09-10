"""Import du dictionnaire Oracle/IFS et génération du modèle SQL report.md.

Seules les métadonnées du catalogue sont stockées. Le SQL produit n'est jamais
exécuté ; les noms Oracle ne sont jamais interpolés dans les requêtes PostgreSQL.

Les fichiers acceptés sont le CSV point-virgule et le classeur Excel .xlsx/.xlsm
(premier onglet, première ligne d'en-têtes) : mêmes en-têtes, mêmes contrôles.
"""
import csv
import io
import json
import re

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.dialects.oracle.base import RESERVED_WORDS

MAX_FILE_BYTES = 32 * 1024 * 1024
XLSX_MAGIC = b'PK\x03\x04'          # .xlsx / .xlsm : archive ZIP
XLS_MAGIC = b'\xd0\xcf\x11\xe0'      # .xls : ancien conteneur OLE2, non lu ici


def _check_headers(headers, required, label, hint):
    if len(set(headers)) != len(headers) or not set(required).issubset(headers):
        raise ValueError(f"{label} : en-têtes requis : {', '.join(required)} ({hint}).")


def _cell(value):
    """Cellule Excel -> texte : 1 et non 1.0 pour un entier, '' pour une cellule vide."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _csv_pairs(content, required, label):
    try:
        decoded = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            decoded = content.decode('cp1252')
        except UnicodeDecodeError as exc:
            raise ValueError(f"{label} : encodage UTF-8 ou Windows-1252 requis.") from exc
    if '\x00' in decoded:
        raise ValueError(f"{label} : contenu CSV invalide.")
    reader = csv.DictReader(io.StringIO(decoded, newline=''), delimiter=';', strict=True)
    try:
        headers = [h.strip() for h in (reader.fieldnames or [])]
        _check_headers(headers, required, label, 'séparateur ;')
        reader.fieldnames = headers
        pairs = []
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"{label}, ligne {reader.line_num} : nombre de cellules incorrect.")
            pairs.append((reader.line_num, row))
        return pairs
    except (csv.Error, UnicodeError) as exc:
        raise ValueError(f"{label} : format CSV invalide.") from exc


def _excel_pairs(content, required, label):
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            table = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    except Exception as exc:  # openpyxl : archive, XML ou onglet invalide
        raise ValueError(f"{label} : classeur Excel illisible (.xlsx ou .xlsm attendu).") from exc
    headers = [value.strip() for value in (table[0] if table else [])]
    while headers and not headers[-1]:  # colonnes vides à droite des en-têtes
        headers.pop()
    _check_headers(headers, required, label, "1re ligne du 1er onglet")
    pairs = []
    for number, row in enumerate(table[1:], start=2):
        if any(value.strip() for value in row[len(headers):]):
            raise ValueError(f"{label}, ligne {number} : nombre de cellules incorrect.")
        values = list(row[:len(headers)]) + [''] * (len(headers) - len(row))
        pairs.append((number, dict(zip(headers, values))))
    return pairs


def _rows(content, required, label):
    """Lignes utiles d'un CSV point-virgule ou d'un classeur Excel, mêmes contrôles."""
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError(f"{label} : fichier vide ou supérieur à 32 Mo.")
    if content.startswith(XLS_MAGIC):
        raise ValueError(f"{label} : ancien format .xls non pris en charge. "
                         "Enregistrez le classeur en .xlsx ou en CSV.")
    reader = _excel_pairs if content.startswith(XLSX_MAGIC) else _csv_pairs
    rows = []
    for number, row in reader(content, required, label):
        row = {key: value.strip() for key, value in row.items()}
        if not any(row.values()):
            continue
        if any(not row[key] for key in required):
            raise ValueError(f"{label}, ligne {number} : valeur obligatoire vide.")
        rows.append(row)
    if not rows:
        raise ValueError(f"{label} : aucune ligne de données.")
    return rows


def _integer(row, key, minimum=None):
    value = row.get(key)
    if not value:
        return None
    try:
        result = int(value)
    except ValueError as exc:
        raise ValueError(f"{row['Table Name']} : {key} doit être un entier.") from exc
    maximum = 9223372036854775807 if key == 'Num Rows' else 2147483647
    if result > maximum or result < -2147483648 or (minimum is not None and result < minimum):
        raise ValueError(f"{row['Table Name']} : valeur hors limites pour {key}.")
    return result


def parse_catalog(tables_csv=None, columns_csv=None):
    """Valide les fichiers fournis entièrement avant toute écriture en base.

    Les deux fichiers sont facultatifs et indépendants : tables seules, colonnes
    seules (leurs tables doivent alors déjà figurer au catalogue) ou les deux.
    Chacun peut être un CSV point-virgule ou un classeur Excel .xlsx/.xlsm.
    """
    table_rows = _rows(tables_csv, ['Owner', 'Table Name'], 'Tables') if tables_csv else []
    column_rows = _rows(columns_csv, ['Owner', 'Table Name', 'Column Name',
                                      'Data Type', 'Nullable', 'Column Id'],
                        'Colonnes') if columns_csv else []
    # Clés naturelles : une ligne répétée écrase la précédente, la dernière gagne.
    # Indispensable aussi côté base : ON CONFLICT ne peut pas traiter deux fois la
    # même clé dans un INSERT groupé.
    tables, columns = {}, {}
    for row in table_rows:
        key = (row['Owner'], row['Table Name'])
        tables[key] = dict(owner=key[0], table_name=key[1],
                           tablespace_name=row.get('Tablespace Name') or None,
                           status=row.get('Status') or None, num_rows=_integer(row, 'Num Rows', 0),
                           metadata=json.dumps(row, ensure_ascii=False))
    for row in column_rows:
        key = (row['Owner'], row['Table Name'])
        column_key = (*key, row['Column Name'])
        # Sans fichier des tables, l'existence est vérifiée en base à l'import.
        if table_rows and key not in tables:
            raise ValueError(f"Table absente du fichier des tables : {'.'.join(key)}.")
        if row['Nullable'] not in ('Y', 'N'):
            raise ValueError(f"{'.'.join(column_key)} : Nullable doit valoir Y ou N.")
        columns[column_key] = dict(owner=key[0], table_name=key[1], column_name=row['Column Name'],
                                   column_id=_integer(row, 'Column Id', 1), data_type=row['Data Type'],
                                   data_length=_integer(row, 'Data Length', 0),
                                   data_precision=_integer(row, 'Data Precision'),
                                   data_scale=_integer(row, 'Data Scale'), nullable=row['Nullable'] == 'Y',
                                   data_default=row.get('Data Default') or None,
                                   metadata=json.dumps(row, ensure_ascii=False))
    # Après écrasement : deux colonnes distinctes ne peuvent pas partager un Column Id.
    positions = set()
    for column in columns.values():
        position = (column['owner'], column['table_name'], column['column_id'])
        if position in positions:
            raise ValueError(f"{column['owner']}.{column['table_name']} : "
                             f"Column Id {column['column_id']} en doublon.")
        positions.add(position)
    return list(tables.values()), list(columns.values())


def import_catalog(engine, tables_csv=None, columns_csv=None):
    if not tables_csv and not columns_csv:
        raise ValueError('Fournissez au moins un fichier : les tables, les colonnes ou les deux.')
    tables, columns = parse_catalog(tables_csv, columns_csv)
    with engine.begin() as connection:
        # Sérialise les imports : les fichiers fournis forment un tout atomique.
        connection.execute(text('SELECT pg_advisory_xact_lock(778813)'))
        if tables:
            connection.execute(text('''
                INSERT INTO public.ifs_table_catalog
                    (owner, table_name, tablespace_name, status, num_rows, metadata)
                VALUES (:owner, :table_name, :tablespace_name, :status, :num_rows, CAST(:metadata AS jsonb))
                ON CONFLICT (owner, table_name) DO UPDATE SET
                    tablespace_name = EXCLUDED.tablespace_name, status = EXCLUDED.status,
                    num_rows = EXCLUDED.num_rows, metadata = EXCLUDED.metadata, imported_at = now()
            '''), tables)
        if not columns:
            return {'tables_imported': len(tables), 'columns_imported': 0}
        ids = {(r.owner, r.table_name): r.table_id for r in connection.execute(text(
            'SELECT table_id, owner, table_name FROM public.ifs_table_catalog'))}
        missing = sorted({(c['owner'], c['table_name']) for c in columns} - set(ids))
        if missing:
            raise ValueError(f"Table absente du catalogue : {'.'.join(missing[0])}. "
                             "Importez d'abord le fichier des tables.")
        for column in columns:
            column['table_id'] = ids[(column.pop('owner'), column.pop('table_name'))]
        query = text('''
            INSERT INTO public.ifs_column_catalog
                (table_id, column_name, column_id, data_type, data_length, data_precision,
                 data_scale, nullable, data_default, metadata)
            VALUES (:table_id, :column_name, :column_id, :data_type, :data_length, :data_precision,
                    :data_scale, :nullable, :data_default, CAST(:metadata AS jsonb))
            ON CONFLICT (table_id, column_name) DO UPDATE SET
                column_id = EXCLUDED.column_id, data_type = EXCLUDED.data_type,
                data_length = EXCLUDED.data_length, data_precision = EXCLUDED.data_precision,
                data_scale = EXCLUDED.data_scale, nullable = EXCLUDED.nullable,
                data_default = EXCLUDED.data_default, metadata = EXCLUDED.metadata, imported_at = now()
        ''')
        for start in range(0, len(columns), 1000):
            connection.execute(query, columns[start:start + 1000])
    return {'tables_imported': len(tables), 'columns_imported': len(columns)}


def build_report(table, columns, selected=None, include_owner=False):
    """Deux SELECT comme docs/ifs_Catalog/report.md, dans l'ordre Column Id."""
    available = {column['column_name'] for column in columns}
    if selected is not None:
        if (not isinstance(selected, list) or not selected
                or any(not isinstance(name, str) for name in selected)
                or len(set(selected)) != len(selected) or not set(selected).issubset(available)):
            raise ValueError('Sélection de colonnes vide, inconnue ou en doublon.')
        available = set(selected)
    ordered = [c['column_name'] for c in sorted(columns, key=lambda c: (c['column_id'], c['column_name']))
               if c['column_name'] in available]
    if not ordered:
        raise ValueError('Cette table ne contient aucune colonne importée.')

    def quoted(name):
        return '"' + name.replace('"', '""') + '"'

    # Identifiants usuels sans guillemets dans le SELECT interne (modèle fourni).
    # Les identifiants atypiques/casse mixte sont échappés, jamais exécutés ici.
    def identifier(name):
        return name if re.fullmatch(r'[A-Z][A-Z0-9_$#]*', name) and name not in RESERVED_WORDS else quoted(name)

    outer = ',\n'.join('    ' + quoted(name) for name in ordered)
    inner = ',\n'.join('        ' + identifier(name) for name in ordered)
    source = identifier(table['table_name'])
    if include_owner:
        source = identifier(table['owner']) + '.' + source
    return f'SELECT\n{outer}\nFROM (\n    SELECT\n{inner}\n    FROM {source});\n'
